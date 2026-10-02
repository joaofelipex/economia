from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = Document()

# Base style
style = doc.styles['Normal']
style.font.name = 'Calibri'
style.font.size = Pt(11)

title = doc.add_heading('Miniteste — DRE, Balanço Patrimonial e DFC', level=0)
subtitle = doc.add_paragraph('Preparação técnica — Itaú BBA')
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_paragraph('Nome: ______________________________          Data: ____ / ____ / ______')
doc.add_paragraph('')

def add_section(titulo):
    doc.add_heading(titulo, level=1)

def add_vf(numero, texto):
    p = doc.add_paragraph()
    p.add_run(f'{numero}. ( V / F )  ').bold = True
    p.add_run(texto)

def add_dissertativa(numero, texto, linhas=4):
    p = doc.add_paragraph()
    p.add_run(f'{numero}. ').bold = True
    p.add_run(texto)
    for _ in range(linhas):
        doc.add_paragraph('_' * 100)
    doc.add_paragraph('')

def add_calculo(numero, texto, linhas=6):
    p = doc.add_paragraph()
    p.add_run(f'{numero}. ').bold = True
    p.add_run(texto)
    for _ in range(linhas):
        doc.add_paragraph('_' * 100)
    doc.add_paragraph('')

# Parte 1
add_section('Parte 1 — Conceitual (Verdadeiro ou Falso)')
add_vf(1, 'O Lucro Bruto é calculado subtraindo as despesas operacionais da Receita Líquida.')
add_vf(2, 'Depreciação é uma despesa que reduz o lucro contábil, mas não representa saída de caixa no período.')
add_vf(3, 'No Balanço Patrimonial, Ativo = Passivo + Patrimônio Líquido sempre deve ser verdadeiro.')
add_vf(4, 'Uma empresa pode ter lucro líquido positivo e mesmo assim apresentar fluxo de caixa operacional negativo.')
add_vf(5, 'Contas a Receber é classificado como Ativo Não Circulante.')
add_vf(6, 'O pagamento de dividendos aparece no Fluxo de Caixa das Atividades de Financiamento.')
add_vf(7, 'EBITDA é calculado como Lucro Líquido + Depreciação/Amortização + Juros + Impostos.')
doc.add_paragraph('')

# Parte 2
add_section('Parte 2 — Conceitual (Dissertativa curta)')
add_dissertativa(8, 'Explique a diferença entre o método direto e o método indireto de construção da DFC.')
add_dissertativa(9, 'Se uma empresa aumenta seu estoque durante o ano (sem alterar vendas), o que acontece com o Fluxo de Caixa Operacional? Por quê?')
add_dissertativa(10, 'Cite três exemplos de itens que ficam no Ativo Circulante e três no Passivo Não Circulante.')
add_dissertativa(11, 'Por que a compra de um imobilizado (ex: uma máquina) não aparece na DRE no momento da compra, mas afeta a DRE nos anos seguintes?')

# Parte 3
add_section('Parte 3 — Prática (Cálculo)')
doc.add_paragraph('Dados de uma empresa fictícia (em R$ mil):')

table = doc.add_table(rows=1, cols=2)
table.style = 'Light Grid Accent 1'
hdr = table.rows[0].cells
hdr[0].text = 'Item'
hdr[1].text = 'Valor'

dados = [
    ('Receita Bruta', '10.000'),
    ('Impostos sobre vendas', '1.500'),
    ('CPV (Custo dos Produtos Vendidos)', '4.000'),
    ('Despesas com Vendas', '800'),
    ('Despesas Administrativas', '600'),
    ('Depreciação (incluída nas despesas acima)', '300'),
    ('Despesas Financeiras', '400'),
    ('Alíquota de Imposto de Renda', '25%'),
]
for item, valor in dados:
    row = table.add_row().cells
    row[0].text = item
    row[1].text = valor

doc.add_paragraph('')

add_calculo(12, 'Monte a DRE completa (Receita Líquida → Lucro Bruto → EBIT → Lucro antes de IR → Lucro Líquido).', linhas=10)
add_calculo(13, 'Calcule o EBITDA dessa empresa.', linhas=4)

doc.add_paragraph('Para a DFC (método indireto), considere também:')
doc.add_paragraph('• Aumento de Contas a Receber: R$ 200 mil', style=None)
doc.add_paragraph('• Redução de Estoques: R$ 100 mil')
doc.add_paragraph('• Aumento de Fornecedores: R$ 150 mil')
doc.add_paragraph('')

add_calculo(14, 'Calcule o Fluxo de Caixa Operacional partindo do Lucro Líquido.', linhas=8)

doc.save(r'C:\Users\Infraestrutura-IMTS\economia\miniteste_dre_balanco_dfc.docx')
print('Documento salvo com sucesso.')
